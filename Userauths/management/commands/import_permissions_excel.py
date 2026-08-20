"""
Commande Django pour importer les permissions personnalisées depuis un fichier Excel.
Format Excel attendu:
- Colonne A: Nom de la permission
- Colonne B: Catégorie
- Colonne C: URL/Name de l'URL
"""
from django.core.management.base import BaseCommand, CommandError
from Userauths.models import CustomPermission, TypeCustomPermission
import os
from openpyxl import load_workbook
from django.db import transaction


class Command(BaseCommand):
    help = 'Importe les permissions personnalisées depuis un fichier Excel'

    def add_arguments(self, parser):
        parser.add_argument(
            'file_path',
            type=str,
            help='Chemin vers le fichier Excel à importer',
        )
        parser.add_argument(
            '--update',
            action='store_true',
            help='Met à jour les permissions existantes au lieu de les ignorer',
        )
        parser.add_argument(
            '--skip-header',
            action='store_true',
            default=True,
            help='Ignore la première ligne (en-têtes)',
        )

    def handle(self, *args, **options):
        file_path = options['file_path']
        update = options['update']
        skip_header = options['skip_header']
        
        if not os.path.exists(file_path):
            raise CommandError(f'❌ Le fichier "{file_path}" n\'existe pas.')
        
        self.stdout.write(self.style.SUCCESS(f'📂 Ouverture du fichier: {file_path}\n'))
        
        try:
            workbook = load_workbook(file_path, data_only=True)
            sheet = workbook.active
            
            self.stdout.write(f'📊 Feuille active: {sheet.title}\n')
            
            created_count = 0
            updated_count = 0
            skipped_count = 0
            error_count = 0
            
            start_row = 2 if skip_header else 1
            
            with transaction.atomic():
                for row_num in range(start_row, sheet.max_row + 1):
                    try:
                        # Lire les colonnes (A=1, B=2, C=3)
                        name = sheet.cell(row=row_num, column=1).value
                        categorie_name = sheet.cell(row=row_num, column=2).value
                        url = sheet.cell(row=row_num, column=3).value
                        
                        # Vérifier que les champs obligatoires sont présents
                        if not name or not url:
                            self.stdout.write(
                                self.style.WARNING(
                                    f'  ⚠️  Ligne {row_num}: Données incomplètes (ignorée)'
                                )
                            )
                            skipped_count += 1
                            continue
                        
                        # Nettoyer les valeurs
                        name = str(name).strip()
                        url = str(url).strip()
                        categorie_name = str(categorie_name).strip() if categorie_name else 'Autres'
                        
                        # Créer ou récupérer la catégorie
                        category, _ = TypeCustomPermission.objects.get_or_create(
                            categorie=categorie_name,
                            defaults={'categorie': categorie_name}
                        )
                        
                        # Créer ou mettre à jour la permission
                        if update:
                            perm, created = CustomPermission.objects.update_or_create(
                                url=url,
                                defaults={
                                    'name': name,
                                    'categorie': category
                                }
                            )
                            if created:
                                created_count += 1
                                self.stdout.write(
                                    self.style.SUCCESS(
                                        f'  ✅ Ligne {row_num}: Créée - {name}'
                                    )
                                )
                            else:
                                updated_count += 1
                                self.stdout.write(
                                    self.style.WARNING(
                                        f'  🔄 Ligne {row_num}: Mise à jour - {name}'
                                    )
                                )
                        else:
                            if CustomPermission.objects.filter(url=url).exists():
                                skipped_count += 1
                                self.stdout.write(
                                    self.style.WARNING(
                                        f'  ⏭️  Ligne {row_num}: Ignorée (existe déjà) - {name}'
                                    )
                                )
                            else:
                                CustomPermission.objects.create(
                                    name=name,
                                    categorie=category,
                                    url=url
                                )
                                created_count += 1
                                self.stdout.write(
                                    self.style.SUCCESS(
                                        f'  ✅ Ligne {row_num}: Créée - {name}'
                                    )
                                )
                    
                    except Exception as e:
                        error_count += 1
                        self.stdout.write(
                            self.style.ERROR(
                                f'  ❌ Ligne {row_num}: Erreur - {str(e)}'
                            )
                        )
            
            # Résumé
            self.stdout.write(self.style.SUCCESS(f'\n✨ Résumé de l\'importation:\n'))
            self.stdout.write(self.style.SUCCESS(f'  ✅ Permissions créées: {created_count}'))
            if update:
                self.stdout.write(self.style.WARNING(f'  🔄 Permissions mises à jour: {updated_count}'))
            if skipped_count > 0:
                self.stdout.write(self.style.WARNING(f'  ⏭️  Permissions ignorées: {skipped_count}'))
            if error_count > 0:
                self.stdout.write(self.style.ERROR(f'  ❌ Erreurs: {error_count}'))
            self.stdout.write(self.style.SUCCESS(f'\n🎉 Importation terminée!'))
        
        except Exception as e:
            raise CommandError(f'❌ Erreur lors de l\'ouverture du fichier Excel: {str(e)}')

