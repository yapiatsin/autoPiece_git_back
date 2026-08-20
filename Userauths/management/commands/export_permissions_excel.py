"""
Commande Django pour exporter les permissions personnalisées vers un fichier Excel.
Format Excel:
- Colonne A: Nom de la permission
- Colonne B: Catégorie
- Colonne C: URL/Name de l'URL
"""
from django.core.management.base import BaseCommand
from Userauths.models import CustomPermission
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import os


class Command(BaseCommand):
    help = 'Exporte les permissions personnalisées vers un fichier Excel'

    def add_arguments(self, parser):
        parser.add_argument(
            '--output',
            type=str,
            default='permissions_export.xlsx',
            help='Nom du fichier de sortie (défaut: permissions_export.xlsx)',
        )
        parser.add_argument(
            '--path',
            type=str,
            default='.',
            help='Chemin où sauvegarder le fichier (défaut: répertoire courant)',
        )

    def handle(self, *args, **options):
        output_name = options['output']
        output_path = options['path']
        
        # Créer le chemin complet
        if not output_name.endswith('.xlsx'):
            output_name += '.xlsx'
        
        full_path = os.path.join(output_path, output_name)
        
        self.stdout.write(self.style.SUCCESS(f'📊 Exportation des permissions vers: {full_path}\n'))
        
        # Récupérer toutes les permissions
        permissions = CustomPermission.objects.all().select_related('categorie').order_by('categorie__categorie', 'name')
        
        if not permissions.exists():
            self.stdout.write(self.style.WARNING('⚠️  Aucune permission à exporter.'))
            return
        
        self.stdout.write(f'📋 {permissions.count()} permissions trouvées\n')
        
        # Créer le classeur Excel
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Permissions'
        
        # Style pour l'en-tête
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        # En-têtes
        headers = ['Nom de la Permission', 'Catégorie', 'URL/Name']
        for col_num, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Largeur des colonnes
        sheet.column_dimensions['A'].width = 40
        sheet.column_dimensions['B'].width = 25
        sheet.column_dimensions['C'].width = 50
        
        # Données
        for row_num, permission in enumerate(permissions, start=2):
            sheet.cell(row=row_num, column=1).value = permission.name
            sheet.cell(row=row_num, column=2).value = permission.categorie.categorie
            sheet.cell(row=row_num, column=3).value = permission.url
        
        # Sauvegarder
        try:
            workbook.save(full_path)
            self.stdout.write(self.style.SUCCESS(f'✅ Fichier créé avec succès: {full_path}'))
            self.stdout.write(self.style.SUCCESS(f'📊 {permissions.count()} permissions exportées'))
            self.stdout.write(self.style.SUCCESS(f'\n🎉 Exportation terminée!'))
        except Exception as e:
            self.stdout.write(self.style.ERROR(f'❌ Erreur lors de la sauvegarde: {str(e)}'))
            raise

